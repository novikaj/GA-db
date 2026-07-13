"""Creates a plain-language companion deck for non-technical stakeholders."""
from pathlib import Path
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).parent
OUT = ROOT / "outputs"
NAVY, BLUE, TEAL, ORANGE, RED, GREY, LIGHT = "132238", "3366CC", "00A6A6", "F59E0B", "DC4C64", "64748B", "F3F6FA"
def rgb(s): return RGBColor.from_string(s)
def money(x): return f"${x/1000:.0f}k"
def pct(x): return f"{x:.1%}"

def rect(s,x,y,w,h,fill=LIGHT,rounded=True):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.color.rgb=rgb(fill); return sh
def text(s,txt,x,y,w,h,size=15,color=NAVY,bold=False):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True; p=tf.paragraphs[0]; p.text=txt; p.font.name="Aptos"; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color); return sh
def title(s,h,sub=""):
    text(s,h,.6,.34,12,.43,25,NAVY,True); text(s,sub,.62,.84,12,.25,10,GREY); rect(s,.6,1.15,1,.04,TEAL,False)
def bullets(s,items,x,y,w,size=14):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(4.8)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if not i else tf.add_paragraph(); p.text=item; p.font.name="Aptos"; p.font.size=Pt(size); p.font.color.rgb=rgb(NAVY); p.space_after=Pt(13)
    return sh
def bar(s,df,cat,val,x,y,w,h,label):
    cd=CategoryChartData(); cd.categories=df[cat].tolist(); cd.add_series(label,df[val].tolist())
    ch=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_legend=False; ch.value_axis.has_major_gridlines=False; ch.category_axis.tick_labels.font.size=Pt(9); ch.value_axis.tick_labels.font.size=Pt(8); ch.series[0].format.fill.solid(); ch.series[0].format.fill.fore_color.rgb=rgb(TEAL); ch.has_title=True; ch.chart_title.text_frame.paragraphs[0].text=label; ch.chart_title.text_frame.paragraphs[0].font.size=Pt(11)

def data():
    ads=pd.read_csv(ROOT/'paid_ads.csv'); crm=pd.read_csv(ROOT/'crm_pipeline.csv'); usage=pd.read_csv(ROOT/'product_usage_events.csv',parse_dates=['event_time'])
    ads['region']=ads.campaign_name.str.extract(r'^(EMEA|APAC|US|LATAM)_')[0]
    crm=crm.merge(ads[['campaign_id','region']].drop_duplicates(),on='campaign_id',how='left'); crm['opportunity']=crm.opportunity_id.notna(); crm['won']=crm.close_status.eq('Won')
    reg=crm.groupby('region',dropna=False).agg(leads=('lead_id','size'), opportunities=('opportunity','sum'), wins=('won','sum'), won_arr=('amount',lambda x:x[crm.loc[x.index,'won']].sum()))
    reg['spend']=ads.groupby('region').cost.sum(); reg['lead_to_win']=reg.wins/reg.leads; reg['roas']=reg.won_arr/reg.spend; reg['cac']=reg.spend/reg.wins
    acct=usage.groupby('company_id').agg(active_days=('event_time',lambda x:x.dt.date.nunique()),events=('event_type','size'))
    accounts=crm.set_index('company_id').join(acct); accounts['activated_30d']=accounts.active_days.ge(30).fillna(False)
    return ads,crm,usage,reg.reset_index(),accounts.reset_index()

def main():
    ads,crm,usage,reg,accounts=data(); prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); b=prs.slide_layouts[6]
    # 1
    s=prs.slides.add_slide(b); rect(s,0,0,13.333,7.5,NAVY,False); text(s,"Growth analytics, explained simply",.78,1.3,11.5,.65,32,"FFFFFF",True); text(s,"What we measure, why it matters, and where growth needs attention",.8,2.15,10.8,.4,18,"C9D7E5"); rect(s,.8,3.05,1.4,.08,TEAL,False); text(s,"A stakeholder-friendly view of paid marketing, sales conversion, and product adoption.",.8,3.45,10.5,.45,20,"FFFFFF")
    # 2
    s=prs.slides.add_slide(b); title(s,"The simple story: people move through a journey","We use the same journey to see where money and customers are being lost.")
    steps=[("1. See an ad","Impressions and clicks\nAre we getting attention?",BLUE),("2. Become a lead","A person shares details\nAre we attracting the right people?",TEAL),("3. Start a trial / sales process","Trial, opportunity\nAre they seriously considering us?",ORANGE),("4. Become a customer","Closed won + ARR\nDid we create revenue?",RED),("5. Keep using product","Logins, active days\nWill value stick?",NAVY)]
    for i,(a,c,col) in enumerate(steps):
        x=.55+i*2.56; rect(s,x,1.9,2.2,2.2,col); text(s,a,x+.15,2.2,1.9,.3,14,"FFFFFF",True); text(s,c,x+.15,2.73,1.9,.58,11,"FFFFFF");
        if i<4: text(s,"→",x+2.22,2.72,.3,.3,22,TEAL,True)
    bullets(s,["Growth analytics is simply the practice of measuring this journey, then improving the weakest step.","We do not judge marketing by clicks alone: the important question is whether it produces customers who get value from the product."],.8,5.05,11.7,16)
    # 3
    s=prs.slides.add_slide(b); title(s,"The five metrics: a small dashboard with a purpose","Together, they answer: Are we buying attention efficiently, converting it, and creating lasting value?")
    items=[("Marketing-sourced leads","How many people came from a measurable campaign?","Tells us whether marketing is creating demand."),("Cost per lead (CPL)","How much ad money did we spend for each lead?","Lets us compare the price of demand."),("Lead → opportunity rate","What share of leads became a serious sales conversation?","Shows lead quality, not just volume."),("CAC and ROAS","Cost to acquire a customer; revenue returned for every $1 spent.","Shows whether growth pays back."),("30-day activation rate","Share of accounts active on 30+ different days.","Early signal that customers found real value." )]
    for i,(a,b1,c) in enumerate(items):
        y=1.4+i*1.03; rect(s,.72,y,12,0.8,"EAF7F7" if i in [0,3] else LIGHT); text(s,a,.95,y+.12,2.75,.25,13,NAVY,True); text(s,b1,3.85,y+.12,4.2,.46,11,BLUE); text(s,c,8.35,y+.12,4.0,.46,11,TEAL)
    #4
    s=prs.slides.add_slide(b); title(s,"Metrics 1–2: are we generating demand at a sensible price?","These are acquisition metrics—the top of the customer journey.")
    rect(s,.75,1.55,5.7,3.85,"EAF7F7"); text(s,"Marketing-sourced leads",1.05,1.9,4.8,.35,19,NAVY,True); text(s,"Plain meaning",1.05,2.53,2,.22,12,TEAL,True); text(s,"People who became leads and can be connected to a marketing campaign.",1.05,2.8,4.8,.5,15,NAVY); text(s,"Why it matters",1.05,3.6,2,.22,12,TEAL,True); text(s,"It measures whether marketing creates potential customers—not merely website traffic.",1.05,3.87,4.8,.52,15,NAVY)
    rect(s,6.85,1.55,5.7,3.85,"FFF3E0"); text(s,"Cost per lead (CPL)",7.15,1.9,4.8,.35,19,NAVY,True); text(s,"Plain meaning",7.15,2.53,2,.22,12,ORANGE,True); text(s,"Ad spend divided by the number of attributable leads.",7.15,2.8,4.8,.5,15,NAVY); text(s,"Why it matters",7.15,3.6,2,.22,12,ORANGE,True); text(s,"A cheap lead is not automatically good; CPL should be read alongside lead quality and revenue.",7.15,3.87,4.8,.52,15,NAVY)
    text(s,"Example: spend $10,000 and receive 20 leads → CPL is $500 per lead.",.9,5.85,11.4,.35,16,TEAL,True)
    #5
    s=prs.slides.add_slide(b); title(s,"Metrics 3–5: do leads become customers who get value?","These metrics prevent us from celebrating cheap—but poor-quality—leads.")
    blocks=[("Lead → opportunity rate","Of every 100 leads, how many become a meaningful sales opportunity?\n\nWhy: tells Sales and Marketing whether the audience is a good fit.",BLUE),("CAC and ROAS","CAC: ad spend ÷ customers won. ROAS: won revenue ÷ ad spend.\n\nWhy: tells Finance whether growth is economically worthwhile.",ORANGE),("30-day activation","Accounts active on at least 30 distinct days.\n\nWhy: regular use is a practical early sign that the product is valuable.",TEAL)]
    for i,(a,b1,col) in enumerate(blocks):
        x=.7+i*4.18; rect(s,x,1.6,3.75,3.85,"F3F6FA"); text(s,a,x+.25,1.93,3.2,.45,18,col,True); text(s,b1,x+.25,2.75,3.15,1.65,14,NAVY)
    text(s,"Why these five? They cover the whole business loop: demand, efficiency, quality, revenue, and customer value—without overwhelming a first-time reader.",.85,6.05,11.5,.45,15,NAVY,True)
    #6 regional
    s=prs.slides.add_slide(b); title(s,"Regional check: EMEA is the clearest area of concern","This comparison uses the region embedded in campaign names; 33 unattributed leads are excluded.")
    order=reg.dropna().sort_values('roas'); bar(s,order,'region','roas',.65,1.5,6.2,4.55,"Revenue returned per $1 of ad spend (ROAS)")
    for i,row in order.reset_index(drop=True).iterrows():
        y=1.62+i*1.18; rect(s,7.35,y,5.2,.92,"FDF2F3" if row.region=='EMEA' else LIGHT); text(s,row.region,7.6,y+.17,1.2,.26,14,NAVY,True); text(s,f"{row.roas:.1f}x ROAS | {pct(row.lead_to_win)} lead-to-win | {int(row.wins)} wins",8.75,y+.17,3.45,.29,12,RED if row.region=='EMEA' else NAVY,True)
    text(s,"EMEA spends nearly as much as APAC ($424k vs. $436k), but produces only 7 wins vs. 21—and just 31¢ of observed won ARR per $1 spent.",.75,6.38,11.8,.4,14,RED,True)
    #7
    s=prs.slides.add_slide(b); title(s,"What is happening in each region?","The broad pattern is consistent; the size of the commercial problem differs.")
    rows=[("EMEA — lagging","0.3× ROAS; 8.8% of leads become wins. Paid Search is particularly weak (0.1× ROAS).","Pause and diagnose EMEA Search: search terms, audience, message, landing page, and speed-to-lead."),("APAC — strong","1.0× ROAS; 17.1% lead-to-win. Display returns 8.6× observed spend.","Protect Display; test incremental budget in controlled steps."),("LATAM — strongest","1.2× ROAS; 16.7% lead-to-win. Display returns 8.3× observed spend.","Use as the benchmark market for campaign and landing-page practices."),("US — middle","1.1× ROAS; best lead-to-opportunity rate (37.7%), but fewer wins than LATAM/APAC.","Investigate sales follow-up and later-stage conversion, not just top-of-funnel volume.")]
    for i,(a,b1,c) in enumerate(rows):
        y=1.4+i*1.25; rect(s,.7,y,12,1.02,"FDF2F3" if i==0 else LIGHT); text(s,a,.95,y+.14,2.2,.35,14,RED if i==0 else NAVY,True); text(s,b1,3.25,y+.13,4.35,.55,11,NAVY); text(s,c,7.85,y+.13,4.4,.55,11,TEAL)
    #8 activation region
    s=prs.slides.add_slide(b); title(s,"One finding is true in every region: sustained usage matters","Product adoption is not a technical detail—it is a leading indicator of sales success.")
    xregions=accounts[accounts.region.notna() & accounts.close_status.eq('Won')].groupby('region').activated_30d.mean().reindex(['EMEA','US','APAC','LATAM']).reset_index(); xregions.columns=['region','activation']
    bar(s,xregions,'region','activation',.65,1.55,6.0,4.25,"30-day activation rate among won accounts")
    bullets(s,["In every measured region, 80%+ of closed-won accounts are active for at least 30 days.","Closed-lost accounts in every region have a median of only 14–15 active days and none reaches 30 active days.","What to do: help trial accounts get to meaningful use quickly—onboarding nudges at day 7, then sales follow-up for engaged accounts at day 14."],7.1,1.75,5.25,15)
    #9
    s=prs.slides.add_slide(b); title(s,"A few data issues need fixing before major budget decisions","The direction is useful today, but these cleanups make the next decision much safer.")
    issues=[("Channel labels disagree","3,003 ad records have a campaign name that says one thing (e.g., Search) and a channel field that says another (e.g., Paid Social).","Use one agreed channel classification, while retaining the original raw fields."),("Plan tier changes in event data","Every account with usage appears under 3–4 different plan tiers.","Keep a separate, dated subscription table; do not infer plan from activity events."),("Different time windows","Ads and CRM end in Sep 2025; product events continue to Jan 2026. 33 leads lack campaign attribution.","Compare mature cohorts only and label unattributed leads Organic/Unknown.")]
    for i,(a,b1,c) in enumerate(issues):
        y=1.42+i*1.5; rect(s,.7,y,12,1.22,"FDF2F3" if i==0 else LIGHT); text(s,a,.95,y+.15,2.6,.3,14,RED if i==0 else NAVY,True); text(s,b1,3.65,y+.15,3.9,.55,11,NAVY); text(s,c,7.9,y+.15,4.2,.55,11,TEAL)
    #10
    s=prs.slides.add_slide(b); title(s,"What I would do next","A practical sequence: fix measurement, improve activation, then scale only validated demand.")
    nexts=[("This month","Agree the campaign taxonomy and dashboard definitions. Set EMEA Paid Search guardrails while investigating the funnel."),(("Next 30–60 days"),"Introduce a trial health score based on active days, logins, and feature use. Trigger onboarding at day 7 and sales outreach at day 14."),("Next 60–90 days","Run a controlled test before increasing Display spend; use a holdout or matched-market test to separate real lift from attribution." )]
    for i,(a,b1) in enumerate(nexts):
        y=1.55+i*1.42; rect(s,.85,y,11.65,1.02,"EAF7F7" if i==0 else LIGHT); text(s,a,1.15,y+.18,2.45,.32,15,TEAL,True); text(s,b1,3.85,y+.18,8.1,.48,14,NAVY)
    text(s,"The goal is not more reporting. It is faster, more confident decisions about where to spend and how to help prospects succeed.",.95,6.2,11.5,.42,16,NAVY,True)
    prs.save(OUT/'growth_analytics_explained.pptx')
    print(OUT/'growth_analytics_explained.pptx')

if __name__=='__main__': main()
